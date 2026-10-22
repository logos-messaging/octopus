"""Build the Octopus hackathon deck (octopus.pptx). Convert with LibreOffice to ODP and PDF.

See slides/README.md.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

INK = RGBColor(0x1B, 0x1F, 0x2A)
MUTED = RGBColor(0x5B, 0x63, 0x75)
ACCENT = RGBColor(0x6D, 0x28, 0xD9)  # octopus purple
ACCENT_SOFT = RGBColor(0xED, 0xE9, 0xFE)
GREEN_SOFT = RGBColor(0xDC, 0xFC, 0xE7)
GREY_SOFT = RGBColor(0xF1, 0xF3, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "DejaVu Sans"
MONO = "DejaVu Sans Mono"

W, H = Emu(12192000), Emu(6858000)  # 16:9
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def inches(x: float) -> Emu:
    return Emu(int(x * 914400))


def text(slide, x, y, w, h, lines, size=20, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=6):
    """A text box. `lines` is a list of str or (str, dict) with per-line overrides."""
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        opts = {}
        if isinstance(line, tuple):
            line, opts = line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("spacing", spacing))
        run = p.add_run()
        run.text = line
        f = run.font
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.name = opts.get("font", font)
        f.color.rgb = opts.get("color", color)
    return box


def _hang(paragraph, left_in, hang_in):
    """A hanging indent: wrapped lines line up with the text, not with the bullet."""
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(int(left_in * 914400)))
    ppr.set("indent", str(-int(hang_in * 914400)))


def bullets(slide, x, y, w, h, items, size=20):
    lines = []
    for item in items:
        if isinstance(item, tuple):  # (text, sub-bullet)
            lines.append(("–\t" + item[0], {"size": size - 3, "color": MUTED, "sub": True}))
        else:
            lines.append(("•\t" + item, {"size": size}))
    tb = text(slide, x, y, w, h, lines, spacing=10)
    for para, line in zip(tb.text_frame.paragraphs, lines):
        if line[1].get("sub"):
            _hang(para, 0.75, 0.3)
        else:
            _hang(para, 0.35, 0.35)
    return tb


def box(slide, x, y, w, h, label, fill=ACCENT_SOFT, size=16, bold=True, color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, inches(x), inches(y), inches(w), inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = ACCENT
    s.line.width = Pt(1.25)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size if i == 0 else size - 3)
        r.font.bold = bold if i == 0 else False
        r.font.name = FONT
        r.font.color.rgb = color if i == 0 else MUTED
    return s


def arrow(slide, x1, y1, x2, y2, label=None, label_dx=0.0, label_dy=-0.35):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inches(x1), inches(y1), inches(x2), inches(y2))
    c.line.color.rgb = ACCENT
    c.line.width = Pt(2)
    line = c.line._get_or_add_ln()
    tail = line.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    line.append(tail)
    if label:
        mx, my = (x1 + x2) / 2 + label_dx, (y1 + y2) / 2 + label_dy
        text(slide, mx - 1.6, my, 3.2, 0.4, [(label, {"size": 12, "color": MUTED, "align": PP_ALIGN.CENTER})])


def header(slide, title, kicker=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if kicker:
        text(slide, 0.6, 0.35, 12, 0.4, [(kicker.upper(), {"size": 12, "color": ACCENT, "bold": True})])
    text(slide, 0.6, 0.65, 12.2, 0.9, [(title, {"size": 32, "bold": True})])


def footer(slide, n):
    text(slide, 0.6, 7.05, 9, 0.35, [("Octopus · Logos Messaging Offsite 2026", {"size": 11, "color": MUTED})])
    text(slide, 12.2, 7.05, 0.8, 0.35, [(str(n), {"size": 11, "color": MUTED, "align": PP_ALIGN.RIGHT})])


def new_slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    header(s, title, kicker)
    footer(s, len(prs.slides))
    return s


# 1. Title -------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
bg.fill.solid()
bg.fill.fore_color.rgb = ACCENT
bg.line.fill.background()
text(s, 0.9, 1.5, 11.5, 1.4, [("Octopus", {"size": 66, "bold": True, "color": WHITE})])
text(s, 0.9, 2.75, 11.5, 0.8, [("Lend your laptop to your team's CI, over Logos Delivery",
                                {"size": 24, "color": ACCENT_SOFT})])
text(s, 0.9, 3.9, 11.5, 1.2, [
    ("Hackathon - Logos Messaging Offsite 2026", {"size": 20, "color": WHITE, "bold": True}),
    ("Logos Delivery", {"size": 18, "color": ACCENT_SOFT}),
])
text(s, 0.9, 5.4, 11.5, 0.8, [("Fabiana Cecin  ·  Tanya Stubbs  ·  Ivan Folgueira Bande",
                               {"size": 20, "color": WHITE})])

# 2. Problem -----------------------------------------------------------------------------------
s = new_slide("CI minutes cost money, and the laptops sit idle", "The problem")
bullets(s, 0.6, 1.7, 7.4, 4.8, [
    "Hosted runners bill by the minute; macOS minutes cost the most",
    "Every developer has a fast machine that is idle most of the day",
    "Self-hosted runners exist, but:",
    ("registering one needs repo admin rights", None),
    ("an always-on coordinator needs a GitHub admin token on a server", None),
    "We wanted: team machines run the CI, with no server and no stored admin token",
])
box(s, 8.5, 1.9, 4.2, 1.1, "GitHub-hosted runner\npaid per minute", fill=GREY_SOFT)
box(s, 8.5, 3.3, 4.2, 1.1, "Team laptop\nidle, already trusted", fill=GREEN_SOFT)
box(s, 8.5, 4.7, 4.2, 1.1, "Coordinator server\nholds an admin token", fill=GREY_SOFT)

# 3. Idea --------------------------------------------------------------------------------------
s = new_slide("Logos Delivery is the control plane", "The idea")
bullets(s, 0.6, 1.7, 12, 3.2, [
    "No coordinator: the team machines and the CI job meet on a Logos Delivery topic",
    "Discovery, job assignment, results and live logs all travel over the network",
    "GitHub only proves who is asking (OIDC token) — it never gets a runner registration",
    "Anything no machine takes, and Windows, falls back to a GitHub-hosted runner",
])
box(s, 0.9, 5.0, 3.4, 1.2, "Discovery\nwho is online, which OS")
box(s, 4.95, 5.0, 3.4, 1.2, "Assignment\nclaim · lease · takeover")
box(s, 9.0, 5.0, 3.4, 1.2, "Results & logs\nencrypted reply channel")

# 4. How it works ------------------------------------------------------------------------------
s = new_slide("How a CI run reaches a laptop", "How it works")
box(s, 0.5, 1.9, 2.6, 1.2, "GitHub\npush / pull request", fill=GREY_SOFT)
box(s, 0.5, 4.3, 2.6, 1.4, "submit job\nhosted, cheap Linux\nOIDC token", fill=GREY_SOFT)
box(s, 4.6, 3.0, 3.6, 1.8, "Logos Delivery\npublic logos.test fleet\ntopic + encrypted channels",
    fill=ACCENT_SOFT, size=18)
box(s, 9.6, 1.5, 3.1, 1.2, "Laptop (macOS)\noctopus ci node", fill=GREEN_SOFT)
box(s, 9.6, 3.3, 3.1, 1.2, "Server (Linux)\noctopus ci node", fill=GREEN_SOFT)
box(s, 9.6, 5.1, 3.1, 1.2, "GitHub-hosted\nfallback · Windows", fill=GREY_SOFT)
arrow(s, 1.8, 3.1, 1.8, 4.3, "runs the workflow", label_dx=1.0, label_dy=-0.1)
arrow(s, 3.1, 4.9, 4.6, 4.2, "request per leg", label_dx=0.2, label_dy=0.3)
arrow(s, 8.2, 3.6, 9.6, 2.1, "claim / lease", label_dy=-0.45)
arrow(s, 8.2, 3.9, 9.6, 3.9)
arrow(s, 8.2, 4.4, 9.6, 5.7, "nobody took it", label_dx=-0.75, label_dy=0.2)
text(s, 0.5, 6.35, 12.3, 0.6, [(
    (
        "Nodes verify the token against GitHub's public keys, rank themselves (rendezvous hash), "
        "claim, run .octopus/ci.sh, and stream the log back. The submit job prints it and gates the PR."
    ),
    {"size": 14, "color": MUTED})])

# 5. Trust -------------------------------------------------------------------------------------
s = new_slide("Who may run what", "Trust and security")
bullets(s, 0.6, 1.7, 6.2, 5.0, [
    "Every request carries a GitHub OIDC token",
    ("checked against GitHub's public keys; one-time use", None),
    ("pinned workflow, allowed repo, no fork PRs", None),
    "Team membership is a signed chain",
    ("team key → member list (GitHub logins, per repo/org)", None),
    ("→ SSH signature of each node, checked on github.com/<login>.keys", None),
    "No admin token on any server",
], size=19)
bullets(s, 6.9, 1.7, 6.0, 5.0, [
    "Jobs can run in a throwaway container",
    ("sees only the checkout; no gh login, SSH keys or home", None),
    ("read-only root, no capabilities, resource limits", None),
    "Honest limits",
    ("a member's machine is trusted to report the truth", None),
    ("container on Linux only; macOS needs a VM", None),
], size=19)

# 6. Demo --------------------------------------------------------------------------------------
s = new_slide("Live demo", "Demo")
steps = [
    ("1", "octopus ci status", "this laptop's proof, team key and allowlist — its only setup"),
    ("2", "git push", "a one-line change starts the workflow on GitHub"),
    ("3", "submit job joins the fleet", "publishes one request per leg on Logos Delivery"),
    ("4", "CI leg started …", "the laptop claims the macOS leg and runs .octopus/ci.sh"),
    ("5", "[macos] …", "its output streams live into the GitHub Actions log"),
    ("6", "job summary", "macOS: this laptop · Linux: server node or hosted · Windows: hosted"),
]
for i, (n, cmd, what) in enumerate(steps):
    y = 1.65 + i * 0.85
    box(s, 0.6, y, 0.6, 0.6, n, fill=ACCENT, color=WHITE, size=18, shape=MSO_SHAPE.OVAL)
    text(s, 1.45, y + 0.02, 4.4, 0.6, [(cmd, {"size": 18, "font": MONO, "bold": True})])
    text(s, 5.9, y + 0.05, 7.0, 0.6, [(what, {"size": 17, "color": MUTED})])

# 7. What Logos Delivery gave us -------------------------------------------------------------
s = new_slide("What Logos Delivery gave us", "Why it works")
cards = [
    ("Reach", "Nodes behind home NATs talk through the public fleet: no ports to open, no server"),
    ("Discovery", "One pub/sub topic: machines announce, CI jobs ask, everyone hears"),
    ("Private channels", "Reliable, end-to-end encrypted channels for results and live logs"),
    ("Fast join", "A fresh node joins the fleet and exchanges messages in about 4 seconds"),
]
for i, (t, d) in enumerate(cards):
    x = 0.6 + (i % 2) * 6.2
    y = 1.8 + (i // 2) * 2.35
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inches(x), inches(y), inches(5.9), inches(2.0))
    card.fill.solid()
    card.fill.fore_color.rgb = ACCENT_SOFT
    card.line.color.rgb = ACCENT
    text(s, x + 0.3, y + 0.2, 5.3, 0.5, [(t, {"size": 22, "bold": True, "color": ACCENT})])
    text(s, x + 0.3, y + 0.8, 5.3, 1.1, [(d, {"size": 17})])

# 8. Dogfooding ------------------------------------------------------------------------------
s = new_slide("We used it on itself", "What we learned")
bullets(s, 0.6, 1.7, 6.3, 5.0, [
    "Octopus runs its own CI through Octopus",
    "We installed it as new users and filed what hurt",
    "Fixed during the hackathon:",
    ("install ~9 min → under a minute (no submodule clones)", None),
    ("library 279 MB → 68 MB (stripped)", None),
    ("TRACE flood ~360 lines/s → INFO by default", None),
    ("stop took 24 s → immediate", None),
    ("clear errors for missing gh login and tokens", None),
], size=19)
bullets(s, 6.9, 1.7, 6.0, 5.0, [
    "Still open:",
    ("a node keeps ≤ 5 peers per IP (library limit)", None),
    ("nightly end-to-end test needs its repo set up", None),
    "Three ways to use it, one CLI:",
    ("octopus ci init — join the team CI", None),
    ("octopus init — a plain runner for a repo", None),
    ("coordinator mode — one-use JIT runners", None),
], size=19)

# 9. Status and next -------------------------------------------------------------------------
s = new_slide("Status and next", "Where it stands")
bullets(s, 0.6, 1.7, 12, 4.8, [
    "v0.4.0 released: pip wheel, prebuilt libraries for Linux x86_64 and macOS arm64",
    "octopus update upgrades a machine and restarts its service",
    "Next:",
    ("“drivers”: GitHub today; GitLab or a plain “run this” later", None),
    ("a macOS VM sandbox, like the Linux container", None),
    ("more team machines, and measuring cost and speed against hosted runners", None),
])

# 10. Thanks ---------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
bg.fill.solid()
bg.fill.fore_color.rgb = ACCENT
bg.line.fill.background()
text(s, 0.9, 2.2, 11.5, 1.2, [("Thank you", {"size": 60, "bold": True, "color": WHITE})])
text(s, 0.9, 3.5, 11.5, 0.8, [("github.com/logos-messaging/octopus", {"size": 24, "color": ACCENT_SOFT,
                                                                      "font": MONO})])
text(s, 0.9, 5.2, 11.5, 0.8, [("Fabiana Cecin  ·  Tanya Stubbs  ·  Ivan Folgueira Bande",
                               {"size": 20, "color": WHITE})])

out = Path(__file__).with_name("octopus.pptx")
prs.save(out)
print("wrote", out, "with", len(prs.slides), "slides")
